import sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
import openpyxl

init_header = "%%{init: {'flowchart': {'htmlLabels': true}, 'themeCSS': '.node foreignObject { overflow: visible !important; } .node foreignObject div { width: 360px !important; text-align: center; line-height: 1.5; white-space: normal; }'} }%%"

target_file = 'GMS_자동진행_시퀀스_표준템플릿_v4.xlsx'
wb = openpyxl.load_workbook(target_file)

# Fix '사용 방법' sheet
ws_intro = wb['사용 방법']
ws_intro.cell(row=52, column=2).value = 'Mermaid.js (flowchart TD)'

# Accurate mapping of sheets and the exact code cell (row, col)
sheet_code_cells = [
    ('예제1_단순개폐', 11, 13, 12, 13),
    ('예제2_조기통과', 11, 13, 12, 13),
    ('예제3_수동확인', 11, 13, 12, 13),
    ('예제4_지연개폐', 10, 13, 11, 13),
    ('예제5_반복루프', 12, 13, 13, 13),
    ('예제6_누출시험', 13, 13, 14, 13),
    ('Bypass_v1', 17, 13, 18, 13),
]

for sheet_name, title_r, title_c, code_r, code_c in sheet_code_cells:
    ws = wb[sheet_name]
    # Restore clean title
    ws.cell(row=title_r, column=title_c).value = '📜 공식 Mermaid 기법 소스 코드 (flowchart TD - 복사용)'
    
    code_val = ws.cell(row=code_r, column=code_c).value
    if code_val and isinstance(code_val, str):
        lines = [l.rstrip() for l in code_val.splitlines()]
        # Remove any init_header lines or markdown backticks
        clean_lines = []
        for l in lines:
            if l.strip().startswith('%%{init:') or l.strip().startswith('```'):
                continue
            if l.strip().startswith('classDef '):
                l = l.replace(',width:320px', '').replace(',width:280px', '').replace(',width:340px', '')
            clean_lines.append(l)
        
        final_code = init_header + '\n' + '\n'.join(clean_lines).strip()
        ws.cell(row=code_r, column=code_c).value = final_code
        print(f'Sheet [{sheet_name}]: Title restored at R{title_r}C{title_c}, Code updated at R{code_r}C{code_c}')

wb.save(target_file)
print('All clean and perfectly saved!')
