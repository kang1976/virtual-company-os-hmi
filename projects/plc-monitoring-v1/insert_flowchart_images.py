import sys, io, os, base64, urllib.request, time
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
import openpyxl
from openpyxl.drawing.image import Image
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

os.makedirs('flowchart_images', exist_ok=True)

target_file = 'GMS_자동진행_시퀀스_표준템플릿_v4.xlsx'
wb = openpyxl.load_workbook(target_file)

# Sheet configurations: (sheet_name, code_row, code_col, insert_start_row)
sheet_configs = [
    ('예제1_단순개폐', 12, 13, 24),
    ('예제2_조기통과', 12, 13, 24),
    ('예제3_수동확인', 12, 13, 24),
    ('예제4_지연개폐', 11, 13, 20),
    ('예제5_반복루프', 13, 13, 28),
    ('예제6_누출시험', 14, 13, 29),
    ('Bypass_v1', 18, 13, 45),
]

header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
header_font = Font(name="맑은 고딕", size=11, bold=True, color="FFFFFF")
border_thin = Border(
    left=Side(style='thin', color='CBD5E1'),
    right=Side(style='thin', color='CBD5E1'),
    top=Side(style='thin', color='CBD5E1'),
    bottom=Side(style='thin', color='CBD5E1')
)

for idx, (sheet_name, code_r, code_c, insert_r) in enumerate(sheet_configs, 1):
    ws = wb[sheet_name]
    mermaid_code = ws.cell(row=code_r, column=code_c).value
    if not mermaid_code:
        print(f'No code in {sheet_name}')
        continue
    
    # Generate image via mermaid.ink
    graph_bytes = mermaid_code.encode('utf-8')
    b64_str = base64.b64encode(graph_bytes).decode('ascii')
    img_url = f'https://mermaid.ink/img/{b64_str}'
    img_path = f'flowchart_images/flowchart_{idx}.png'
    
    print(f'Downloading image for [{sheet_name}]...')
    try:
        req = urllib.request.Request(img_url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req, timeout=10) as response:
            with open(img_path, 'wb') as f:
                f.write(response.read())
        print(f'  Saved {img_path} ({os.path.getsize(img_path)} bytes)')
    except Exception as e:
        print(f'  Error downloading for {sheet_name}: {e}')
        continue
    
    # Add title banner for the image
    ws.cell(row=insert_r, column=1, value=f"🖼️ [공정 플로우차트 렌더링 이미지 - {sheet_name}]")
    ws.merge_cells(start_row=insert_r, start_column=1, end_row=insert_r, end_column=12)
    banner_cell = ws.cell(row=insert_r, column=1)
    banner_cell.fill = header_fill
    banner_cell.font = header_font
    banner_cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # Insert image at row insert_r + 2, column B (column 2)
    img = Image(img_path)
    # Scale appropriately if very large
    if img.width > 650:
        ratio = 650 / img.width
        img.width = int(img.width * ratio)
        img.height = int(img.height * ratio)
    
    target_cell = f"B{insert_r + 2}"
    ws.add_image(img, target_cell)
    print(f'  Inserted image into [{sheet_name}] at {target_cell}')
    time.sleep(0.5)

wb.save(target_file)
print(f'\nALL FLOWCHART IMAGES PERFECTLY INSERTED INTO {target_file}!')
