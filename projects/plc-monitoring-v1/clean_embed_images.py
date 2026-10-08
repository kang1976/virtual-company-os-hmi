import sys, io, os
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
import openpyxl
from openpyxl.drawing.image import Image as OpenpyxlImage
from openpyxl.styles import Font, PatternFill, Alignment

target_file = 'GMS_자동진행_시퀀스_표준템플릿_v4.xlsx'
wb = openpyxl.load_workbook(target_file)

diagrams = [
    ('예제1_단순개폐', 24, 1),
    ('예제2_조기통과', 24, 2),
    ('예제3_수동확인', 24, 3),
    ('예제4_지연개폐', 20, 4),
    ('예제5_반복루프', 28, 5),
    ('예제6_누출시험', 29, 6),
    ('Bypass_v1', 45, 7)
]

header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
header_font = Font(name="맑은 고딕", size=11, bold=True, color="FFFFFF")

for sheet_name, insert_r, idx in diagrams:
    ws = wb[sheet_name]
    # Clear any old images in the sheet
    ws._images.clear()
    
    # Title banner
    ws.cell(row=insert_r, column=1, value=f"🖼️ [공정 플로우차트 와이드(Wide) 고화질 다이어그램 - {sheet_name}]")
    ws.merge_cells(start_row=insert_r, start_column=1, end_row=insert_r, end_column=12)
    banner_cell = ws.cell(row=insert_r, column=1)
    banner_cell.fill = header_fill
    banner_cell.font = header_font
    banner_cell.alignment = Alignment(horizontal="center", vertical="center")
    
    # Add exactly one fresh HD wide image
    img_path = f'flowchart_images_hd/flowchart_{idx}.png'
    img = OpenpyxlImage(img_path)
    target_cell = f"B{insert_r + 2}"
    ws.add_image(img, target_cell)
    print(f"Cleanly embedded 1 image into [{sheet_name}] at {target_cell}")

wb.save(target_file)
print("All sheets now have exactly 1 clean HD wide flowchart image!")
