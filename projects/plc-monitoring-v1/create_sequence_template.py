import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()

# 스타일 공통 정의
font_title = Font(name="맑은 고딕", size=15, bold=True, color="FFFFFF")
font_section = Font(name="맑은 고딕", size=12, bold=True, color="1E3A8A")
font_warn_section = Font(name="맑은 고딕", size=12, bold=True, color="991B1B")
font_header = Font(name="맑은 고딕", size=10, bold=True, color="FFFFFF")
font_bold = Font(name="맑은 고딕", size=10, bold=True)
font_normal = Font(name="맑은 고딕", size=10)
font_link = Font(name="맑은 고딕", size=10, bold=True, color="1D4ED8", underline="single")

fill_title = PatternFill(start_color="1E40AF", end_color="1E40AF", fill_type="solid")
fill_section = PatternFill(start_color="DBEAFE", end_color="DBEAFE", fill_type="solid")
fill_warn_section = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
fill_hdr_blue = PatternFill(start_color="1D4ED8", end_color="1D4ED8", fill_type="solid")
fill_hdr_red = PatternFill(start_color="B91C1C", end_color="B91C1C", fill_type="solid")
fill_hdr_amber = PatternFill(start_color="B45309", end_color="B45309", fill_type="solid")
fill_hdr_green = PatternFill(start_color="047857", end_color="047857", fill_type="solid")
fill_zebra = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
fill_highlight = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
fill_green_tag = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")

thin_border = Border(
    left=Side(style='thin', color='CBD5E1'),
    right=Side(style='thin', color='CBD5E1'),
    top=Side(style='thin', color='CBD5E1'),
    bottom=Side(style='thin', color='CBD5E1')
)

align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
align_left = Alignment(horizontal="left", vertical="center", wrap_text=True)
align_right = Alignment(horizontal="right", vertical="center")

# 공통 컬럼 정의
fixed_cols = [
    "S/No.", "Main Step", "Sub Step", "Next Step", "Operations", "Cycle",
    "진행방식", "Ack Goto", "Alarm Goto", "Message at Controller", "Time (Sec)", "Acc,Time (Sec)"
]
valve_cols = [
    "VN1", "VN2", "V/S_{side}", "AG_{side}", "PNV", "GNV", "HPIV",
    "FPV_{side}", "LPI_{side}", "HPV_{side}", "HPI_{side}", "LPV_{side}", "PGII_{side}"
]
trailing_cols = [
    "Alarm Monitoring", "비교연산자", "설정명(비교대상ID)", "조기통과", "Alarm Seq.", "Alarm Message", "Remarks"
]
all_headers = fixed_cols + valve_cols + trailing_cols
total_cols = len(all_headers)
valve_start = 13
valve_end = 12 + len(valve_cols)
trail_start = valve_end + 1
trail_end = total_cols


# ========================================================
# 1. 메인 시트: "사용 방법" (하이퍼링크 목록 포함)
# ========================================================
ws_guide = wb.active
ws_guide.title = "사용 방법"
ws_guide.views.sheetView[0].showGridLines = True

# 타이틀 배너
ws_guide.merge_cells("A1:G1")
ws_guide["A1"] = "📘 GMS 자동진행 시퀀스 엑셀 마스터 매뉴얼 & 상황별 바로가기"
ws_guide["A1"].font = font_title
ws_guide["A1"].fill = fill_title
ws_guide["A1"].alignment = align_center
ws_guide.row_dimensions[1].height = 42

ws_guide.merge_cells("A2:G2")
ws_guide["A2"] = "본 시트는 GMS 시퀀스 규칙 설명과 함께, 아래 5개 상황별 실제 엑셀 시트로 즉시 이동할 수 있는 바로가기 링크를 제공합니다."
ws_guide["A2"].font = font_bold
ws_guide["A2"].alignment = align_left
ws_guide.row_dimensions[2].height = 24

# 섹션 0: 글씨 깨짐 방지 규칙
ws_guide.merge_cells("A4:G4")
ws_guide["A4"] = "★ [필독] 글씨 깨짐(PK... 외계어) 원인 및 올바른 파일 오픈 규칙"
ws_guide["A4"].font = font_warn_section
ws_guide["A4"].fill = fill_warn_section
ws_guide.row_dimensions[4].height = 28

for col_idx, h in enumerate(["구분", "발생 원인", "화면에 나타나는 증상", "해결책 및 올바른 사용 방법"], 1):
    cell = ws_guide.cell(row=5, column=col_idx, value=h)
    cell.font = font_header
    cell.fill = fill_hdr_red
    cell.alignment = align_center
ws_guide.merge_cells("D5:G5")
ws_guide.row_dimensions[5].height = 26

crash_rules = [
    ("원인 1: 뷰어 오사용", "에디터/메모장으로 열기", "PK... docProps/app.xml 등 외계어 및 글자 깨짐 발생", "★ .xlsx는 압축 바이너리 파일입니다. 메모장이나 VS Code 기본 텍스트 뷰어로 열지 마시고, 반드시 [Microsoft Excel], [한셀], [LibreOffice] 등 정식 스프레드시트 프로그램으로 열어야 합니다."),
    ("원인 2: 확장자 오류", ".csv를 .xlsx로 강제 변경", "서버 업로드 시 '시트를 찾을 수 없습니다' 파싱 오류", "CSV 파일을 이름만 .xlsx로 바꾸면 안 됩니다. 반드시 엑셀에서 [다른 이름으로 저장 ➔ Excel 통합 문서(*.xlsx)]로 정식 저장하세요."),
    ("원인 3: 한글 인코딩", "UTF-8 인코딩 손실", "웹 화면에 물음표 다이아몬드 깨짐", "엑셀 내 한글 셀 내용(공정명, 메시지, 밸브명)은 표준 UTF-8 인코딩을 따르며, 저장 시 기본 통합 문서 포맷을 유지하면 100% 정상 인식됩니다.")
]

for r_idx, cr in enumerate(crash_rules, 6):
    ws_guide.cell(row=r_idx, column=1, value=cr[0]).alignment = align_center
    ws_guide.cell(row=r_idx, column=2, value=cr[1]).alignment = align_center
    ws_guide.cell(row=r_idx, column=3, value=cr[2]).alignment = align_left
    ws_guide.merge_cells(f"D{r_idx}:G{r_idx}")
    ws_guide.cell(row=r_idx, column=4, value=cr[3]).alignment = align_left
    ws_guide.row_dimensions[r_idx].height = 28
    for c in range(1, 8):
        ws_guide.cell(row=r_idx, column=c).border = thin_border
        ws_guide.cell(row=r_idx, column=c).font = font_normal

# 섹션 1: 5가지 상황별 엑셀 예제 시트 바로가기 (★ 핵심 요청 사항: 링크 연결)
ws_guide.merge_cells("A10:G10")
ws_guide["A10"] = "1. 실무 5대 상황별 엑셀 실제 예제 시트 바로가기 (클릭 시 해당 시트로 이동)"
ws_guide["A10"].font = font_section
ws_guide["A10"].fill = fill_section
ws_guide.row_dimensions[10].height = 28

for col_idx, h in enumerate(["상황 구분", "예제 시트 바로가기 (링크)", "상황 명칭 및 핵심 동작", "초보자 핵심 학습 포인트"], 1):
    cell = ws_guide.cell(row=11, column=col_idx, value=h)
    cell.font = font_header
    cell.fill = fill_hdr_amber
    cell.alignment = align_center
ws_guide.merge_cells("C11:D11")
ws_guide.merge_cells("E11:G11")
ws_guide.row_dimensions[11].height = 26

scenarios = [
    ("상황 1", "예제1_단순개폐", "단순 밸브 개폐 및 안정화 대기 (상태 유지)", "필요한 밸브만 O/C 입력하고, 빈칸(공백)은 이전 상태를 유지하는 방법 학습"),
    ("상황 2", "예제2_조기통과", "진공 배기 조기통과 (Early Pass)", "120초 설정했으나 0.5 Torr 도달 시 5초 만에 즉시 다음 스텝 통과하는 방법 학습"),
    ("상황 3", "예제3_수동확인", "작업자 육안 확인 대기 (수동 ACK)", "진행방식 '확인' 설정 시 작업자가 버튼을 누를 때까지 정지 대기하는 방법 학습"),
    ("상황 4", "예제4_지연개폐", "밸브 순차 지연 개폐 (시차 제어)", "O+2, O+4 문법을 사용하여 밸브를 2초, 4초 간격으로 순차 개폐하는 방법 학습"),
    ("상황 5", "예제5_반복루프", "다회 반복 퍼지 루프 (Cycle 3회)", "'진행횟수'와 Next Step을 이용하여 목표 횟수만큼 자동 반복 루프를 도는 방법 학습"),
    ("종합본", "Bypass_v1", "5가지 상황이 모두 합쳐진 종합 실전 시퀀스", "실제 현장 가스 캐비닛 장비에 즉시 업로드 가능한 완전한 시퀀스 양식 학습")
]

for r_idx, sc in enumerate(scenarios, 12):
    ws_guide.cell(row=r_idx, column=1, value=sc[0]).alignment = align_center
    
    # 하이퍼링크 생성
    linkCell = ws_guide.cell(row=r_idx, column=2, value=f"👉 [{sc[1]}] 시트 열기")
    linkCell.hyperlink = f"#'{sc[1]}'!A1"
    linkCell.font = font_link
    linkCell.alignment = align_center
    
    ws_guide.merge_cells(f"C{r_idx}:D{r_idx}")
    ws_guide.cell(row=r_idx, column=3, value=sc[2]).alignment = align_left
    
    ws_guide.merge_cells(f"E{r_idx}:G{r_idx}")
    ws_guide.cell(row=r_idx, column=5, value=sc[3]).alignment = align_left
    
    ws_guide.row_dimensions[r_idx].height = 26
    for c in range(1, 8):
        ws_guide.cell(row=r_idx, column=c).border = thin_border
        if c != 2:
            ws_guide.cell(row=r_idx, column=c).font = font_normal

# 섹션 2: 엑셀 시트 작성 3대 절대 원칙
ws_guide.merge_cells("A19:G19")
ws_guide["A19"] = "2. 엑셀 시트 작성 3대 절대 원칙 (구조 규칙)"
ws_guide["A19"].font = font_section
ws_guide["A19"].fill = fill_section
ws_guide.row_dimensions[19].height = 28

for col_idx, h in enumerate(["항목", "규칙명", "핵심 규칙 및 시스템 동작 원리"], 1):
    cell = ws_guide.cell(row=20, column=col_idx, value=h)
    cell.font = font_header
    cell.fill = fill_hdr_blue
    cell.alignment = align_center
ws_guide.merge_cells("C20:E20")
ws_guide.merge_cells("F20:G20")
ws_guide.cell(row=20, column=6, value="주의사항 (절대 금지)").font = font_header
ws_guide.cell(row=20, column=6).fill = fill_hdr_blue
ws_guide.cell(row=20, column=6).alignment = align_center

rules = [
    ("원칙 1", "시트 이름 규칙", "엑셀의 워크시트 이름이 시스템의 서브시퀀스 ID가 됩니다 (예: Bypass_v1, OneP_v1).", "시트명을 임의로 바꾸면 시스템이 다른 시퀀스로 인식하거나 누락됩니다."),
    ("원칙 2", "행(Row) 위치 고정", "Row 1~3은 타이틀/안내 영역, ★Row 4가 프로그램 기준 헤더 행, Row 5부터 실제 스텝 데이터입니다.", "Row 4의 열 이름을 변경하거나 헤더 위치를 위/아래로 옮기면 프로그램이 읽지 못합니다."),
    ("원칙 3", "상태 유지(빈칸) 원칙", "밸브 셀에 아무것도 적지 않은 공백(빈칸)은 '이전 스텝의 개폐 상태를 그대로 유지'함을 의미합니다.", "매 스텝마다 모든 밸브를 O/C로 채울 필요가 없으며, 빈칸이 통신 부하와 장비 쇼크를 막습니다.")
]

for r_idx, r in enumerate(rules, 21):
    ws_guide.cell(row=r_idx, column=1, value=r[0]).alignment = align_center
    ws_guide.cell(row=r_idx, column=2, value=r[1]).alignment = align_center
    ws_guide.merge_cells(f"C{r_idx}:E{r_idx}")
    ws_guide.merge_cells(f"F{r_idx}:G{r_idx}")
    ws_guide.cell(row=r_idx, column=3, value=r[2]).alignment = align_left
    ws_guide.cell(row=r_idx, column=6, value=r[3]).alignment = align_left
    ws_guide.row_dimensions[r_idx].height = 24
    for c in range(1, 8):
        ws_guide.cell(row=r_idx, column=c).border = thin_border
        ws_guide.cell(row=r_idx, column=c).font = font_normal

ws_guide.column_dimensions['A'].width = 16
ws_guide.column_dimensions['B'].width = 24
ws_guide.column_dimensions['C'].width = 14
ws_guide.column_dimensions['D'].width = 28
ws_guide.column_dimensions['E'].width = 46
ws_guide.column_dimensions['F'].width = 24
ws_guide.column_dimensions['G'].width = 16


# ========================================================
# 공통 템플릿 빌더 함수
# ========================================================
def build_scenario_sheet(sheet_title, banner_title, scenario_desc, steps):
    ws = wb.create_sheet(title=sheet_title)
    ws.views.sheetView[0].showGridLines = True
    
    # 1행: 타이틀
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_cols)
    ws.cell(row=1, column=1, value=f"{sheet_title} - {banner_title}").font = font_title
    ws.cell(row=1, column=1).fill = fill_title
    ws.cell(row=1, column=1).alignment = align_center
    ws.row_dimensions[1].height = 36
    
    # 2행: 메인 설명 및 '사용 방법'으로 돌아가기 링크
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=12)
    back_cell = ws.cell(row=2, column=1, value="🔙 [사용 방법] 매뉴얼 시트로 돌아가기")
    back_cell.hyperlink = "#'사용 방법'!A1"
    back_cell.font = font_link
    back_cell.alignment = align_center
    back_cell.fill = fill_section
    
    ws.merge_cells(start_row=2, start_column=13, end_row=2, end_column=total_cols)
    ws.cell(row=2, column=13, value=f"💡 상황 설명: {scenario_desc}").font = font_bold
    ws.cell(row=2, column=13).alignment = align_left
    ws.cell(row=2, column=13).fill = fill_highlight
    ws.row_dimensions[2].height = 24
    
    # 3행: 카테고리 헤더
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=12)
    ws.cell(row=3, column=1, value="[ 공정 제어 및 타이머 설정 구간 ]").font = font_header
    ws.cell(row=3, column=1).fill = fill_hdr_blue
    ws.cell(row=3, column=1).alignment = align_center
    
    ws.merge_cells(start_row=3, start_column=valve_start, end_row=3, end_column=valve_end)
    ws.cell(row=3, column=valve_start, value="[ 밸브 개폐 제어 구간 (O:열림 / C:닫힘 / 빈칸:유지) ]").font = font_header
    ws.cell(row=3, column=valve_start).fill = fill_hdr_amber
    ws.cell(row=3, column=valve_start).alignment = align_center
    
    ws.merge_cells(start_row=3, start_column=trail_start, end_row=3, end_column=trail_end)
    ws.cell(row=3, column=trail_start, value="[ 센서 감시 / 조기통과 / 알람 판정 구간 ]").font = font_header
    ws.cell(row=3, column=trail_start).fill = fill_hdr_green
    ws.cell(row=3, column=trail_start).alignment = align_center
    ws.row_dimensions[3].height = 20
    
    # 4행: 실제 헤더 행
    ws.row_dimensions[4].height = 28
    for c_idx, h_text in enumerate(all_headers, 1):
        c = ws.cell(row=4, column=c_idx, value=h_text)
        c.font = font_header
        c.alignment = align_center
        c.border = thin_border
        if c_idx <= 12:
            c.fill = PatternFill(start_color="1D4ED8", end_color="1D4ED8", fill_type="solid")
        elif c_idx <= valve_end:
            c.fill = PatternFill(start_color="B45309", end_color="B45309", fill_type="solid")
        else:
            c.fill = PatternFill(start_color="047857", end_color="047857", fill_type="solid")
            
    # 5행부터 스텝 데이터
    for row_offset, step in enumerate(steps, 5):
        ws.row_dimensions[row_offset].height = 24
        row_data = [
            step["no"], step["mainStep"], step["subStep"], step["nextStep"],
            step["op"], step["cycle"], step["adv"], step["ackGoto"],
            step["alarmGoto"], step["msg"], step["time"], f"=SUM($K$5:K{row_offset})"
        ]
        for v_tag in valve_cols:
            row_data.append(step["valves"].get(v_tag, ""))
        row_data.extend([
            step["mon"], step["opCond"], step["valCond"], step["early"],
            step["alarmSeq"], step["alarmMsg"], step["rem"]
        ])
        
        for c_idx, val in enumerate(row_data, 1):
            cell = ws.cell(row=row_offset, column=c_idx, value=val)
            cell.font = font_normal
            cell.border = thin_border
            
            if c_idx in [1, 2, 3, 4, 6, 7, 8, 9]:
                cell.alignment = align_center
            elif c_idx in [5, 10]:
                cell.alignment = align_left
            elif c_idx in [11, 12]:
                cell.alignment = align_right
            elif c_idx <= valve_end:
                cell.alignment = align_center
                if val in ["O", "C"] or str(val).startswith("O+") or str(val).startswith("C+"):
                    cell.font = Font(name="맑은 고딕", size=10, bold=True, color="1E3A8A")
                    cell.fill = fill_highlight
            else:
                if c_idx == trail_start + 3:
                    cell.alignment = align_center
                    if val == "Y":
                        cell.font = Font(name="맑은 고딕", size=10, bold=True, color="047857")
                        cell.fill = fill_green_tag
                else:
                    cell.alignment = align_left
                    
    # 열 너비 조정
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        max_len = max(len(str(cell.value or '')) for cell in col)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)
    ws.column_dimensions['E'].width = 30
    ws.column_dimensions['J'].width = 36
    ws.column_dimensions['L'].width = 16
    ws.column_dimensions[get_column_letter(total_cols)].width = 44


# ========================================================
# 2. 예제 1 시트: "예제1_단순개폐"
# ========================================================
steps_ex1 = [
    {"no": 1, "mainStep": 1, "subStep": 1, "nextStep": "", "op": "[ 전 밸브 안전 Close ]", "cycle": "", "adv": "자동", "ackGoto": "", "alarmGoto": "", "msg": "공정 시작 전 전 밸브 Close", "time": 5, "valves": {"VN1": "C", "VN2": "C", "PNV": "C", "HPV_{side}": "C"}, "mon": "", "opCond": "", "valCond": "", "early": "", "alarmSeq": "", "alarmMsg": "", "rem": "모든 밸브를 안전하게 닫고 5초 대기"},
    {"no": 2, "mainStep": 1, "subStep": 2, "nextStep": "", "op": "[ 질소 가스 공급 개방 ]", "cycle": "", "adv": "자동", "ackGoto": "", "alarmGoto": "", "msg": "질소 공급 밸브 개방 유지", "time": 30, "valves": {"VN1": "O", "PNV": "O"}, "mon": "", "opCond": "", "valCond": "", "early": "", "alarmSeq": "", "alarmMsg": "", "rem": "★핵심: VN1, PNV만 O 입력하고 나머지 밸브는 빈칸으로 상태 유지!"},
    {"no": 3, "mainStep": 1, "subStep": 3, "nextStep": "", "op": "[ 질소 차단 및 대기 ]", "cycle": "", "adv": "자동", "ackGoto": "", "alarmGoto": "", "msg": "질소 공급 밸브 차단", "time": 10, "valves": {"VN1": "C", "PNV": "C"}, "mon": "", "opCond": "", "valCond": "", "early": "", "alarmSeq": "", "alarmMsg": "", "rem": "공급 밸브 닫고 10초간 안정화 대기"}
]
build_scenario_sheet("예제1_단순개폐", "단순 밸브 개폐 및 안정화 대기", "원하는 밸브만 O/C 입력하고, 빈칸으로 두면 이전 개폐 상태를 그대로 유지합니다.", steps_ex1)


# ========================================================
# 3. 예제 2 시트: "예제2_조기통과"
# ========================================================
steps_ex2 = [
    {"no": 1, "mainStep": 2, "subStep": 1, "nextStep": "", "op": "[ 진공 라인 밸브 개방 ]", "cycle": "", "adv": "자동", "ackGoto": "", "alarmGoto": "", "msg": "진공 배기 밸브 Open", "time": 5, "valves": {"HPV_{side}": "O"}, "mon": "", "opCond": "", "valCond": "", "early": "", "alarmSeq": "", "alarmMsg": "", "rem": "진공 배기 라인 개방"},
    {"no": 2, "mainStep": 2, "subStep": 2, "nextStep": "", "op": "[ 고진공 배기 (조기통과 적용) ]", "cycle": "", "adv": "자동", "ackGoto": "", "alarmGoto": 99, "msg": "목표 진공 도달 시 조기 통과", "time": 120, "valves": {}, "mon": "VPT_{side}", "opCond": "<=", "valCond": "0.5", "early": "Y", "alarmSeq": 1, "alarmMsg": "120초 내 진공 미도달 (누설 의심)", "rem": "★핵심: 120초 설정했으나 0.5 Torr 도달 시 5초 만에 즉시 다음 스텝 패스!"},
    {"no": 3, "mainStep": 2, "subStep": 3, "nextStep": "", "op": "[ 진공 유지 및 다음 공정 ]", "cycle": "", "adv": "자동", "ackGoto": "", "alarmGoto": "", "msg": "진공 도달 완료 후 다음 단계 진행", "time": 10, "valves": {"HPV_{side}": "C"}, "mon": "", "opCond": "", "valCond": "", "early": "", "alarmSeq": "", "alarmMsg": "", "rem": "조기통과 후 즉시 진공 밸브 Close"}
]
build_scenario_sheet("예제2_조기통과", "진공 배기 조기통과 (Early Pass)", "조기통과 'Y'를 입력하면 설정 시간(120초)이 다 안 지나도 목표 진공(0.5 Torr) 도달 즉시 패스합니다.", steps_ex2)


# ========================================================
# 4. 예제 3 시트: "예제3_수동확인"
# ========================================================
steps_ex3 = [
    {"no": 1, "mainStep": 3, "subStep": 1, "nextStep": "", "op": "[ 시험 압력 가압 진행 ]", "cycle": "", "adv": "자동", "ackGoto": "", "alarmGoto": "", "msg": "시험 가스 주입 중", "time": 20, "valves": {"VN1": "O", "PNV": "O"}, "mon": "", "opCond": "", "valCond": "", "early": "", "alarmSeq": "", "alarmMsg": "", "rem": "시험 압력까지 가스 주입"},
    {"no": 2, "mainStep": 3, "subStep": 2, "nextStep": "", "op": "[ 작업자 육안 확인 대기 (ACK) ]", "cycle": "", "adv": "확인", "ackGoto": "", "alarmGoto": "", "msg": "압력 게이지 지침 확인 후 화면의 [확인] 버튼을 누르세요", "time": 0, "valves": {"VN1": "C", "PNV": "C"}, "mon": "", "opCond": "", "valCond": "", "early": "", "alarmSeq": "", "alarmMsg": "", "rem": "★핵심: 진행방식에 '확인'을 적으면 작업자가 확인 버튼을 누를 때까지 정지 대기"},
    {"no": 3, "mainStep": 3, "subStep": 3, "nextStep": "", "op": "[ 가압 해제 및 완료 ]", "cycle": "", "adv": "자동", "ackGoto": "", "alarmGoto": "", "msg": "작업자 확인 완료 후 배기 진행", "time": 10, "valves": {"HPV_{side}": "O"}, "mon": "", "opCond": "", "valCond": "", "early": "", "alarmSeq": "", "alarmMsg": "", "rem": "확인 후 자동 진행"}
]
build_scenario_sheet("예제3_수동확인", "작업자 육안 확인 대기 (수동 ACK)", "진행방식 열에 '확인'(또는 ack)을 적으면 작업자가 화면의 [확인] 버튼을 누를 때까지 안전 대기합니다.", steps_ex3)


# ========================================================
# 5. 예제 4 시트: "예제4_지연개폐"
# ========================================================
steps_ex4 = [
    {"no": 1, "mainStep": 4, "subStep": 1, "nextStep": "", "op": "[ 배관 충격 방지 순차 지연 개폐 ]", "cycle": "", "adv": "자동", "ackGoto": "", "alarmGoto": "", "msg": "밸브가 2초, 4초 간격으로 순차 개방됩니다", "time": 15, "valves": {"VN2": "O", "LPV_{side}": "O+2", "HPIV": "O+4"}, "mon": "", "opCond": "", "valCond": "", "early": "", "alarmSeq": "", "alarmMsg": "", "rem": "★핵심: VN2는 즉시 열림, LPV는 2초 뒤, HPIV는 4초 뒤 순차적으로 열림!"},
    {"no": 2, "mainStep": 4, "subStep": 2, "nextStep": "", "op": "[ 순차 지연 닫힘 제어 ]", "cycle": "", "adv": "자동", "ackGoto": "", "alarmGoto": "", "msg": "밸브가 순차적으로 닫힙니다", "time": 10, "valves": {"VN2": "C", "LPV_{side}": "C+2", "HPIV": "C+3"}, "mon": "", "opCond": "", "valCond": "", "early": "", "alarmSeq": "", "alarmMsg": "", "rem": "★핵심: C+n 문법으로 닫힘 동작도 시차 제어 가능"}
]
build_scenario_sheet("예제4_지연개폐", "밸브 순차 지연 개폐 (시차 제어)", "O+n 또는 C+n 문법(예: O+2, C+3)을 사용하여 밸브가 시간차를 두고 순차 동작하도록 제어합니다.", steps_ex4)


# ========================================================
# 6. 예제 5 시트: "예제5_반복루프"
# ========================================================
steps_ex5 = [
    {"no": 1, "mainStep": 5, "subStep": 1, "nextStep": "", "op": "[ 루프 준비 - 밸브 초기화 ]", "cycle": "", "adv": "자동", "ackGoto": "", "alarmGoto": "", "msg": "퍼지 루프 시작 전 준비", "time": 5, "valves": {"VN1": "C", "PNV": "C", "HPV_{side}": "C"}, "mon": "", "opCond": "", "valCond": "", "early": "", "alarmSeq": "", "alarmMsg": "", "rem": "반복 퍼지 준비"},
    {"no": 2, "mainStep": 5, "subStep": 2, "nextStep": "", "op": "[ 퍼지 가스 가압 (루프 시작점) ]", "cycle": "", "adv": "자동", "ackGoto": "", "alarmGoto": "", "msg": "질소 가스 가압 주입 중", "time": 10, "valves": {"VN1": "O", "PNV": "O", "HPV_{side}": "C"}, "mon": "", "opCond": "", "valCond": "", "early": "", "alarmSeq": "", "alarmMsg": "", "rem": "반복 루프 시작 스텝"},
    {"no": 3, "mainStep": 5, "subStep": 3, "nextStep": 2, "op": "[ 배기 진행 및 3회 반복 판정 ]", "cycle": "", "adv": "자동", "ackGoto": "", "alarmGoto": "", "msg": "배기 진행 중 (설정 3회 반복)", "time": 15, "valves": {"VN1": "C", "PNV": "C", "HPV_{side}": "O"}, "mon": "진행횟수", "opCond": "<", "valCond": "3", "early": "", "alarmSeq": "", "alarmMsg": "", "rem": "★핵심: 진행횟수가 3회 미만이면 Next Step 2로 점프! 3회 완료 시 4번 스텝 진행!"},
    {"no": 4, "mainStep": 5, "subStep": 4, "nextStep": "", "op": "[ 3회 반복 완료 및 정상 종료 ]", "cycle": "", "adv": "자동", "ackGoto": "", "alarmGoto": "", "msg": "3회 퍼지 사이클 정상 완료", "time": 5, "valves": {"HPV_{side}": "C"}, "mon": "", "opCond": "", "valCond": "", "early": "", "alarmSeq": "", "alarmMsg": "", "rem": "루프 완료 후 전 밸브 Close"}
]
build_scenario_sheet("예제5_반복루프", "다회 반복 퍼지 루프 (Cycle 3회)", "Alarm Monitoring에 '진행횟수'를 적고 Next Step을 루프 시작 번호로 지정하여 3회 반복 사이클을 수행합니다.", steps_ex5)


# ========================================================
# 7. 종합본 시트: "Bypass_v1" (실전 전체 양식)
# ========================================================
steps_bypass = [
    {"no": 1, "mainStep": 11, "subStep": 1, "nextStep": "", "op": "[ 준비 - 전 밸브 초기화 ]", "cycle": "", "adv": "자동", "ackGoto": "", "alarmGoto": "", "msg": "Bypass 공정 시작 - 전 밸브 안전 초기화", "time": 5, "valves": {"VN1": "C", "VN2": "C", "PNV": "C", "GNV": "C", "HPV_{side}": "C", "LPV_{side}": "C", "HPIV": "C"}, "mon": "", "opCond": "", "valCond": "", "early": "", "alarmSeq": "", "alarmMsg": "", "rem": "★상황 1: 전 밸브 안전 Close 후 5초 대기"},
    {"no": 2, "mainStep": 11, "subStep": 2, "nextStep": "", "op": "[ 1차 질소 공급 및 퍼지 ]", "cycle": "", "adv": "자동", "ackGoto": "", "alarmGoto": "", "msg": "질소 공급 밸브 개방 (안정화 대기)", "time": 20, "valves": {"VN1": "O", "PNV": "O"}, "mon": "", "opCond": "", "valCond": "", "early": "", "alarmSeq": "", "alarmMsg": "", "rem": "★상황 1: VN1, PNV만 열고 나머지 밸브는 빈칸으로 상태 유지"},
    {"no": 3, "mainStep": 11, "subStep": 3, "nextStep": "", "op": "[ 2차 진공 배기 (조기통과) ]", "cycle": "", "adv": "자동", "ackGoto": "", "alarmGoto": 99, "msg": "배관 진공 배기 진행 중", "time": 120, "valves": {"VN1": "C", "PNV": "C", "HPV_{side}": "O"}, "mon": "VPT_{side}", "opCond": "<=", "valCond": "0.5", "early": "Y", "alarmSeq": 1, "alarmMsg": "120초 내 진공 미도달 (누설 의심)", "rem": "★상황 2: 120초 전이라도 0.5 Torr 도달 시 즉시 패스!"},
    {"no": 4, "mainStep": 11, "subStep": 4, "nextStep": "", "op": "[ 작업자 육안 확인 (ACK 대기) ]", "cycle": "", "adv": "확인", "ackGoto": "", "alarmGoto": "", "msg": "게이지 누설 확인 후 화면의 [확인] 버튼을 누르세요", "time": 0, "valves": {}, "mon": "", "opCond": "", "valCond": "", "early": "", "alarmSeq": "", "alarmMsg": "", "rem": "★상황 3: 작업자가 버튼 누를 때까지 무한 대기"},
    {"no": 5, "mainStep": 11, "subStep": 5, "nextStep": "", "op": "[ 밸브 순차 지연 개폐 제어 ]", "cycle": "", "adv": "자동", "ackGoto": "", "alarmGoto": "", "msg": "충격 방지 밸브 순차 개방 중", "time": 15, "valves": {"VN2": "O", "LPV_{side}": "O+2", "HPIV": "O+4"}, "mon": "", "opCond": "", "valCond": "", "early": "", "alarmSeq": "", "alarmMsg": "", "rem": "★상황 4: VN2 즉시 열림, LPV 2초 뒤, HPIV 4초 뒤 순차 열림"},
    {"no": 6, "mainStep": 11, "subStep": 6, "nextStep": "", "op": "[ 3회 반복 퍼지 루프 - 가스 주입 ]", "cycle": "", "adv": "자동", "ackGoto": "", "alarmGoto": "", "msg": "퍼지 가스 가압 진행", "time": 10, "valves": {"VN1": "O", "PNV": "O", "HPV_{side}": "C"}, "mon": "", "opCond": "", "valCond": "", "early": "", "alarmSeq": "", "alarmMsg": "", "rem": "★상황 5: 3회 반복 사이클 시작 (가압)"},
    {"no": 7, "mainStep": 11, "subStep": 7, "nextStep": 6, "op": "[ 3회 반복 퍼지 루프 - 배기 및 판정 ]", "cycle": "", "adv": "자동", "ackGoto": "", "alarmGoto": "", "msg": "배기 진행 및 반복 횟수 판정", "time": 15, "valves": {"VN1": "C", "PNV": "C", "HPV_{side}": "O"}, "mon": "진행횟수", "opCond": "<", "valCond": "3", "early": "", "alarmSeq": "", "alarmMsg": "", "rem": "★상황 5: 3회 미만이면 Next Step 6으로 점프, 3회 만족 시 8번 진행"},
    {"no": 8, "mainStep": 11, "subStep": 8, "nextStep": "", "op": "[ 공정 완료 및 전 밸브 안전 Close ]", "cycle": "", "adv": "자동", "ackGoto": "", "alarmGoto": "", "msg": "Bypass 전 공정 정상 완료되었습니다", "time": 5, "valves": {"VN1": "C", "VN2": "C", "PNV": "C", "HPV_{side}": "C", "LPV_{side}": "C", "HPIV": "C"}, "mon": "", "opCond": "", "valCond": "", "early": "", "alarmSeq": "", "alarmMsg": "", "rem": "전 공정 완료 후 안전 닫힘"},
    {"no": 99, "mainStep": 99, "subStep": 99, "nextStep": "", "op": "[ 비상 알람 조치 스텝 ]", "cycle": "", "adv": "확인", "ackGoto": "", "alarmGoto": "", "msg": "이상 발생으로 비상 정지되었습니다. 설비를 점검하세요.", "time": 0, "valves": {"VN1": "C", "VN2": "C", "PNV": "C", "GNV": "C", "HPV_{side}": "C", "LPV_{side}": "C"}, "mon": "", "opCond": "", "valCond": "", "early": "", "alarmSeq": 1, "alarmMsg": "비상 알람 발생 (안전 Close 완료)", "rem": "Step 3에서 알람 시 점프하는 비상 대응 스텝"}
]
build_scenario_sheet("Bypass_v1", "5대 실전 상황 종합 예제", "5가지 상황이 모두 들어간 완성형 종합 실전 시퀀스입니다. 웹 화면에 즉시 업로드 가능합니다.", steps_bypass)

# 파일 저장 (버전 명시 파일명으로 저장하여 잠금 방지)
output_path = r"d:\AI_Work\Antigravity\06.CEO\projects\plc-monitoring-v1\GMS_자동진행_시퀀스_표준템플릿_v2.xlsx"
wb.save(output_path)
# 이전 파일 덮어쓰기 시도 (닫혀있으면 함께 갱신)
try:
    import shutil
    shutil.copyfile(output_path, r"d:\AI_Work\Antigravity\06.CEO\projects\plc-monitoring-v1\GMS_자동진행_시퀀스_템플릿.xlsx")
except Exception:
    pass
print("상황별 독립 시트 및 하이퍼링크 포함 템플릿 생성 완료:", output_path)
